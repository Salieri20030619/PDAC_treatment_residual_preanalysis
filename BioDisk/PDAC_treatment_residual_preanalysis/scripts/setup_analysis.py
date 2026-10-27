from pathlib import Path
import requests,json,subprocess,sys,hashlib,datetime,importlib.util
ROOT=Path('/mnt/BioDisk/wangmingyuan/PDAC_treatment_residual_preanalysis')
PROXY='http://127.0.0.1:7897'; proxies={'http':PROXY,'https':PROXY}
tasks=[]
links=json.loads((ROOT/'00_docs/Werba2023.html.supplement_links.json').read_text())
for suffix in ['MOESM1_ESM.pdf','MOESM5_ESM.xlsx']:
    url=next(x for x in links if suffix in x); tasks.append((url.split('/')[-1],url))
tasks.append(('Hwang2022_PMC.html','https://pmc.ncbi.nlm.nih.gov/articles/PMC10290535/'))
manifest=[]
for name,url in tasks:
    dest=ROOT/'00_docs'/name
    if not dest.exists():
        r=requests.get(url,proxies=proxies,timeout=(30,300));r.raise_for_status();dest.write_bytes(r.content)
    manifest.append({'url':url,'file':str(dest),'bytes':dest.stat().st_size,'sha256':hashlib.sha256(dest.read_bytes()).hexdigest(),'utc':datetime.datetime.now(datetime.timezone.utc).isoformat()})
    if name.endswith('.pdf'):
        assert dest.read_bytes()[:4]==b'%PDF'
        p=subprocess.run(['pdftotext','-layout',str(dest),str(dest.with_suffix('.txt'))]);p.check_returncode()
    if name.endswith('.xlsx'):
        import openpyxl
        wb=openpyxl.load_workbook(dest,read_only=True,data_only=True)
        data={w.title:[[x if x is not None else '' for x in row] for row in w.iter_rows(values_only=True)] for w in wb}
        (ROOT/'01_metadata/Werba2023_genesets.json').write_text(json.dumps(data,indent=2))
        print('GENESET_SHEETS',[(w.title,w.max_row,w.max_column) for w in wb],flush=True)
(ROOT/'01_metadata/additional_reference_manifest.json').write_text(json.dumps(manifest,indent=2))
vendor=ROOT/'scripts/vendor';vendor.mkdir(exist_ok=True)
packages=['infercnvpy==0.6.1']
for package in ['pybiomart','pyreadr','pytoml','pytest_benchmark']:
    if importlib.util.find_spec(package) is None: packages.append(package.replace('_','-'))
command=[sys.executable,'-m','pip','install','--no-deps','--target',str(vendor),'--proxy',PROXY]+packages
print('RUN_COMMAND',command,flush=True)
p=subprocess.run(command);p.check_returncode()
sys.path.insert(0,str(vendor))
import infercnvpy
print('INFERCNVPY_IMPORT_PASS',infercnvpy.__version__,flush=True)
print('ANALYSIS_SETUP_COMPLETE',flush=True)
