from pathlib import Path
import re, json, requests, hashlib, datetime
from html.parser import HTMLParser
ROOT=Path('/mnt/BioDisk/wangmingyuan/PDAC_treatment_residual_preanalysis')
PROXY={'http':'http://127.0.0.1:7897','https':'http://127.0.0.1:7897'}
class Plain(HTMLParser):
    def __init__(self): super().__init__(); self.lines=[]; self.skip=0
    def handle_starttag(self,tag,attrs):
        if tag in ['script','style']: self.skip+=1
        if tag in ['p','h1','h2','h3','h4','tr','li']: self.lines.append('\n')
    def handle_endtag(self,tag):
        if tag in ['script','style']: self.skip=max(0,self.skip-1)
        if tag in ['p','h1','h2','h3','h4','tr','li']: self.lines.append('\n')
    def handle_data(self,data):
        if not self.skip: self.lines.append(data)
manifest=[]
for study,target in [('Werba2023','41467_2023_36296_MOESM3_ESM.xlsx'),('Hwang2022','41588_2022_1134_MOESM4_ESM.xlsx')]:
    html=(ROOT/'00_docs'/f'{study}.html').read_text(); parser=Plain(); parser.feed(html)
    lines=[' '.join(x.split()) for x in ''.join(parser.lines).splitlines() if x.strip()]
    (ROOT/'00_docs'/f'{study}_readable.txt').write_text('\n'.join(lines))
    print('\nSTUDY',study)
    for i,line in enumerate(lines):
        if any(x in line.lower() for x in ['doublet','infercnv','copy number','cnv','data availability','losartan','nivolumab','quality control','500 detectable']):
            print(i,line[:2500])
    links=json.loads((ROOT/'00_docs'/f'{study}.html.supplement_links.json').read_text())
    url=next(x for x in links if target in x)
    dest=ROOT/'00_docs'/target
    if not dest.exists():
        response=requests.get(url,proxies=PROXY,timeout=(30,180)); response.raise_for_status(); dest.write_bytes(response.content)
    assert dest.read_bytes()[:2]==b'PK'
    manifest.append({'file':str(dest),'url':url,'bytes':dest.stat().st_size,'sha256':hashlib.sha256(dest.read_bytes()).hexdigest(),'utc':datetime.datetime.now(datetime.timezone.utc).isoformat()})
    import openpyxl
    wb=openpyxl.load_workbook(dest,read_only=True,data_only=True)
    print('SHEETS',[(w.title,w.max_row,w.max_column) for w in wb])
    sheets={w.title:[[str(x) if x is not None else '' for x in row] for row in w.iter_rows(values_only=True)] for w in wb}
    (ROOT/'01_metadata'/f'{study}_supplement_sheets.json').write_text(json.dumps(sheets,indent=2))
    for key,rows in sheets.items():
        print('SHEET_PREVIEW',key,rows[:8])
(ROOT/'01_metadata/clinical_download_manifest.json').write_text(json.dumps(manifest,indent=2))
print('CLINICAL_SOURCE_FETCH_COMPLETE')
